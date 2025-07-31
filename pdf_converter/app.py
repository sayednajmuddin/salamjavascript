import sys
import os
import webbrowser
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QLabel, QPushButton,
    QFileDialog, QScrollArea, QComboBox, QCheckBox, QLineEdit,
    QProgressBar, QMessageBox, QSplitter
)
from PyQt5.QtGui import QPixmap, QImage
from PyQt5.QtCore import Qt
import fitz  # PyMuPDF
from pdf2docx import Converter
import pytesseract
from pdf2image import convert_from_path
from docx import Document
from PIL import Image
from PIL import ImageOps
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


LANGUAGES = {
    "English": {
        "open_pdf": "\U0001F4C2 Open PDF",
        "zoom_in": "\U0001F50D Zoom In",
        "zoom_out": "\U0001F50E Zoom Out",
        "convert_word": "\U0001F4DD Convert with Engin 1",
        "convert_py": "\U0001F6E0 Convert with Engin 2",
        "start": "Start Page",
        "end": "End Page",
        "status": "\U0001F4C4 No PDF loaded",
        "converted": "\u2705 Converted successfully",
        "no_file": "Please load a PDF file first.",
        "ocr": "\U0001F9E0 Use OCR if needed",
        "dark_mode": "\U0001F319 Dark Mode",
        "show_in_folder": "\U0001F4C2 Show in Folder",
        "success": "\U0001F389 Success",
        "file_saved": "Your file was saved here:",
        "language": "\U0001F310 Language"
    },
    "Dari": {
        "open_pdf": "\U0001F4C2 \u0628\u0627\u0632 \u06A9\u0631\u062F\u0646 PDF",
        "zoom_in": "\U0001F50D \u0632\u0648\u0645 \u062F\u0627\u062E\u0644",
        "zoom_out": "\U0001F50E \u0632\u0648\u0645 \u0628\u06CC\u0631\u0648\u0646",
        "convert_word": "\U0001F4DD \u062A\u0628\u062F\u06CC\u0644 \u0628\u0627 Engin 1",
        "convert_py": "\U0001F6E0 \u062A\u0628\u062F\u06CC\u0644 \u0628\u0627 Engin 2",
        "start": "\u0635\u0641\u062D\u0647 \u0634\u0631\u0648\u0639",
        "end": "\u0635\u0641\u062D\u0647 \u067E\u0627\u06CC\u0627\u0646",
        "status": "\U0001F4C4 \u0647\u06CC\u0686 \u0641\u0627\u06CC\u0644 PDF \u0628\u0627\u0631\u06AF\u0630\u0627\u0631\u06CC \u0646\u0634\u062F\u0647",
        "converted": "\u2705 \u0645\u0648\u0641\u0642\u0627\u0646\u0647 \u062A\u0628\u062F\u06CC\u0644 \u0634\u062F",
        "no_file": "\u0644\u0637\u0641\u0627\u064B \u06CC\u06A9 \u0641\u0627\u06CC\u0644 PDF \u0627\u0646\u062A\u062E\u0627\u0628 \u06A9\u0646\u06CC\u062F.",
        "ocr": "\U0001F9E0 \u0627\u0633\u062A\u0641\u0627\u062F\u0647 \u0627\u0632 OCR \u062F\u0631 \u0635\u0648\u0631\u062A \u0646\u06CC\u0627\u0632",
        "dark_mode": "\U0001F319 \u062D\u0627\u0644\u062A \u062A\u0627\u0631\u06CC\u06A9",
        "show_in_folder": "\U0001F4C2 \u0628\u0627\u0632 \u06A9\u0631\u062F\u0646 \u062F\u0631 \u067E\u0648\u0634\u0647",
        "success": "\U0001F389 \u0645\u0648\u0641\u0642\u06CC\u062A",
        "file_saved": "\u0641\u0627\u06CC\u0644 \u0634\u0645\u0627 \u0627\u06CC\u0646\u062C\u0627 \u0630\u062E\u06CC\u0631\u0647 \u0634\u062F:",
        "language": "\U0001F310 \u0632\u0628\u0627\u0646"
    },
    "Pashto": {
        "open_pdf": "\U0001F4C2 \u062F PDF \u062E\u0644\u0627\u0635\u0648\u0644",
        "zoom_in": "\U0001F50D \u0632\u0648\u0645 \u062F\u0646\u0646\u0647",
        "zoom_out": "\U0001F50E \u0632\u0648\u0645 \u0628\u0647\u0631",
        "convert_word": "\U0001F4DD \u062F Engin 1 \u0633\u0631\u0647 \u062A\u0628\u062F\u06CC\u0644\u0648\u0644",
        "convert_py": "\U0001F6E0 \u062F Engin 2 \u0633\u0631\u0647 \u062A\u0628\u062F\u06CC\u0644\u0648\u0644",
        "start": "\u062F \u067E\u06D0\u0644 \u067E\u0627\u069A\u0647",
        "end": "\u062F \u067E\u0627\u06CC \u067E\u0627\u069A\u0647",
        "status": "\U0001F4C4 \u0647\u06CC\u0686 PDF \u0646\u062F\u064A \u067E\u0648\u0631\u062A\u0647 \u0634\u0648\u06CC",
        "converted": "\u2705 \u0628\u0631\u06CC\u0627\u0644\u064A \u062A\u0628\u062F\u06CC\u0644",
        "no_file": "\u0645\u0647\u0631\u0628\u0627\u0646\u064A \u0648\u06A9\u0693\u0626 PDF \u0641\u0627\u06CC\u0644 \u0627\u0646\u062A\u062E\u0627\u0628 \u06A9\u0693\u0626.",
        "ocr": "\U0001F9E0 \u06A9\u0647 \u0627\u0693\u062A\u06CC\u0627 \u0648\u064A OCR \u0648\u06A9\u0627\u0631\u0648\u0626",
        "dark_mode": "\U0001F319 \u062A\u06CC\u0627\u0631\u0647 \u062D\u0627\u0644\u062A",
        "show_in_folder": "\U0001F4C2 \u067E\u0647 \u0641\u0648\u0644\u0689\u0631 \u06A9\u06D0 \u069A\u0648\u062F\u0644",
        "success": "\U0001F389 \u0628\u0631\u06CC\u0627",
        "file_saved": "\u0633\u062A\u0627\u0633\u0648 \u0641\u0627\u06CC\u0644 \u062F\u0644\u062A\u0647 \u062E\u0648\u0646\u062F\u064A \u0634\u0648:",
        "language": "\U0001F310 \u069A\u0628\u0647"
    }
}

DARK_STYLE = """
    QWidget {
        background-color: #2b2b2b;
        color: #f0f0f0;
    }
    QPushButton {
        background-color: #444;
        color: white;
        border: 1px solid #666;
        border-radius: 8px;
        padding: 10px;
        font-weight: bold;
    }
    QLineEdit, QComboBox {
        background-color: #555;
        color: white;
        border: 1px solid #777;
        border-radius: 6px;
    }
    QScrollArea {
        background-color: #2b2b2b;
    }
"""


class PDFViewer(QWidget):
    def __init__(self):
        super().__init__()
        self.zoom_level = 1.0
        self.current_path = None
        self.layout = QVBoxLayout()
        self.scroll = QScrollArea()
        self.inner = QWidget()
        self.inner_layout = QVBoxLayout(self.inner)
        self.placeholder = QLabel("\U0001F4C2 Drag & Drop your PDF here", alignment=Qt.AlignCenter)
        self.placeholder.setStyleSheet(
            "font-size: 22px; color: #888; border: 2px dashed #aaa; margin: 50px; padding: 50px;"
        )
        self.inner_layout.addWidget(self.placeholder)
        self.scroll.setWidget(self.inner)
        self.scroll.setWidgetResizable(True)
        self.layout.addWidget(self.scroll)
        self.setLayout(self.layout)
        self.setAcceptDrops(True)

    def load_pdf(self, path):
        self.current_path = path
        while self.inner_layout.count():
            widget = self.inner_layout.takeAt(0).widget()
            if widget:
                widget.deleteLater()
        try:
            doc = fitz.open(path)
            viewport_width = self.scroll.viewport().width() - 40
            for page in doc:
                scale = (viewport_width / page.rect.width) * self.zoom_level
                pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
                fmt = QImage.Format_RGBA8888 if pix.alpha else QImage.Format_RGB888
                img = QImage(pix.samples, pix.width, pix.height, pix.stride, fmt)
                lbl = QLabel()
                lbl.setPixmap(QPixmap.fromImage(img))
                lbl.setAlignment(Qt.AlignCenter)
                lbl.setStyleSheet(
                    "padding: 10px; background-color: #ccc; border-radius: 10px;"
                )
                self.inner_layout.addWidget(lbl)
            doc.close()
        except Exception as e:
            print("Error loading PDF:", e)

    def resizeEvent(self, event):
        if self.current_path:
            self.load_pdf(self.current_path)
        super().resizeEvent(event)

    def zoom_in(self):
        self.zoom_level += 0.1
        if self.current_path:
            self.load_pdf(self.current_path)

    def zoom_out(self):
        if self.zoom_level > 0.2:
            self.zoom_level -= 0.1
            if self.current_path:
                self.load_pdf(self.current_path)

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        for url in e.mimeData().urls():
            path = url.toLocalFile()
            if path.endswith(".pdf"):
                self.load_pdf(path)
                self.parent().pdf_path = path


class MainApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("PDF Converter for Pashto and Dari Languages")
        self.lang = "English"
        self.pdf_path = None
        screen = QApplication.primaryScreen().availableGeometry()
        self.resize(int(screen.width() * 0.8), int(screen.height() * 0.8))
        self.center_on_screen()
        self.viewer = PDFViewer()
        self.control = self.build_controls()
        splitter = QSplitter(Qt.Horizontal)
        splitter.addWidget(self.viewer)
        splitter.addWidget(self.control)
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 1)
        self.setCentralWidget(splitter)
        self.update_texts()

    def build_controls(self):
        w = QWidget()
        l = QVBoxLayout()
        self.label_status = QLabel()
        self.box_lang = QComboBox()
        self.box_lang.addItems(LANGUAGES.keys())
        self.box_lang.currentTextChanged.connect(self.set_language)
        self.btn_load = QPushButton()
        self.btn_load.clicked.connect(self.load_file)
        self.btn_zoom_in = QPushButton()
        self.btn_zoom_in.clicked.connect(self.viewer.zoom_in)
        self.btn_zoom_out = QPushButton()
        self.btn_zoom_out.clicked.connect(self.viewer.zoom_out)
        self.chk_ocr = QCheckBox()
        self.chk_dark = QCheckBox()
        self.chk_dark.stateChanged.connect(self.toggle_dark)
        self.input_start = QLineEdit()
        self.input_end = QLineEdit()
        self.btn_word = QPushButton()
        self.btn_word.clicked.connect(self.convert_via_word)
        self.btn_py = QPushButton()
        self.btn_py.clicked.connect(self.convert_via_py)
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        for btn in [self.btn_load, self.btn_zoom_in, self.btn_zoom_out, self.btn_word, self.btn_py]:
            btn.setStyleSheet("padding: 12px; border-radius: 8px; font-weight: bold;")
        l.setSpacing(10)
        for widget in [
            self.label_status,
            self.box_lang,
            self.btn_load,
            self.btn_zoom_in,
            self.btn_zoom_out,
            QLabel("Start:"),
            self.input_start,
            QLabel("End:"),
            self.input_end,
            self.chk_ocr,
            self.chk_dark,
            self.btn_word,
            self.btn_py,
            self.progress,
        ]:
            l.addWidget(widget)
        l.addStretch()
        w.setLayout(l)
        return w

    def center_on_screen(self):
        frameGm = self.frameGeometry()
        center = QApplication.primaryScreen().availableGeometry().center()
        frameGm.moveCenter(center)
        self.move(frameGm.topLeft())

    def update_texts(self):
        t = LANGUAGES[self.lang]
        self.btn_load.setText(t["open_pdf"])
        self.btn_zoom_in.setText(t["zoom_in"])
        self.btn_zoom_out.setText(t["zoom_out"])
        self.label_status.setText(t["status"])
        self.btn_word.setText(t["convert_word"])
        self.btn_py.setText(t["convert_py"])
        self.chk_ocr.setText(t["ocr"])
        self.chk_dark.setText(t["dark_mode"])

    def set_language(self, lang):
        self.lang = lang
        self.update_texts()

    def toggle_dark(self, state):
        self.setStyleSheet(DARK_STYLE if state else "")

    def load_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open PDF", "", "PDF Files (*.pdf)")
        if path:
            self.pdf_path = path
            self.label_status.setText(f"{LANGUAGES[self.lang]['status']}: {os.path.basename(path)}")
            self.viewer.load_pdf(path)

    def convert_via_word(self):
        if not self.pdf_path:
            QMessageBox.warning(self, "No PDF", LANGUAGES[self.lang]["no_file"])
            return
        save, _ = QFileDialog.getSaveFileName(self, "Save Word", "", "Word Files (*.docx)")
        if not save:
            return
        if not save.endswith(".docx"):
            save += ".docx"
        try:
            self.progress.setVisible(True)
            self.progress.setValue(10)
            import win32com.client

            word = win32com.client.Dispatch("Word.Application")
            doc = word.Documents.Open(os.path.abspath(self.pdf_path))
            doc.SaveAs(save, FileFormat=16)
            doc.Close()
            word.Quit()
            self.progress.setValue(100)
            self.success_msg(save)
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))
        finally:
            self.progress.setVisible(False)

    def convert_via_py(self):
        if not self.pdf_path:
            QMessageBox.warning(self, "No PDF", LANGUAGES[self.lang]["no_file"])
            return

        save, _ = QFileDialog.getSaveFileName(self, "Save Word", "", "Word Files (*.docx)")
        if not save:
            return
        if not save.endswith(".docx"):
            save += ".docx"
        try:
            self.progress.setVisible(True)
            self.progress.setValue(10)
            start = (
                int(self.input_start.text()) if self.input_start.text().isdigit() else 1
            )
            end = (
                int(self.input_end.text()) if self.input_end.text().isdigit() else None
            )
            if self.chk_ocr.isChecked():
                images = convert_from_path(
                    self.pdf_path, dpi=400, first_page=start, last_page=end
                )
                ocr_lang = "fas+pus"
                print(f"\u2699 OCR language: {ocr_lang}")
                doc = Document()

                def set_rtl(paragraph, size=12, style="Normal"):
                    paragraph.style = style
                    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    run.font.name = "XB Zar"
                    run.font.size = Pt(size)
                    rPr = run._element.get_or_add_rPr()
                    rtl = OxmlElement("w:rtl")
                    rtl.set(qn("w:val"), "1")
                    rPr.append(rtl)
                    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                    pPr = paragraph._p.get_or_add_pPr()
                    bidi = OxmlElement("w:bidi")
                    bidi.set(qn("w:val"), "1")
                    pPr.append(bidi)
                    jc = OxmlElement("w:jc")
                    jc.set(qn("w:val"), "right")
                    pPr.append(jc)

                for page_index, img in enumerate(images):
                    img = img.convert("L")
                    img = ImageOps.autocontrast(img)
                    text = pytesseract.image_to_string(
                        img, lang=ocr_lang, config="--oem 3 --psm 6"
                    )
                    self.progress.setValue(
                        int((page_index + 1) / len(images) * 80)
                    )
                    lines = [line.strip() for line in text.split("\n") if line.strip()]
                    for line_index, line in enumerate(lines):
                        para = doc.add_paragraph(line)
                        if page_index == 0 and line_index == 0:
                            set_rtl(para, size=16, style="Heading1")
                        elif page_index == 0 and line_index == 1:
                            set_rtl(para, size=14, style="Heading2")
                        else:
                            set_rtl(para, size=12)
                    if page_index < len(images) - 1:
                        doc.add_page_break()
                doc.save(save)
            else:
                conv = Converter(self.pdf_path)
                conv.convert(save, start=start, end=end)
                conv.close()
            self.progress.setValue(100)
            self.success_msg(save)
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))
        finally:
            self.progress.setVisible(False)

    def success_msg(self, path):
        msg = QMessageBox(self)
        msg.setWindowTitle(LANGUAGES[self.lang]["success"])
        msg.setText(f"{LANGUAGES[self.lang]['file_saved']} {path}")
        btn = msg.addButton(LANGUAGES[self.lang]["show_in_folder"], QMessageBox.ActionRole)
        msg.addButton("OK", QMessageBox.AcceptRole)
        msg.exec_()
        if msg.clickedButton() == btn:
            webbrowser.open(os.path.dirname(path))


if __name__ == "__main__":
    app = QApplication(sys.argv)
    main = MainApp()
    main.show()
    sys.exit(app.exec_())
